"""
DOCX Generation Service using python-docx
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from typing import Dict
import logging
from io import BytesIO

logger = logging.getLogger(__name__)


class DOCXService:
    """Service for generating DOCX resumes"""

    def generate_docx(self, resume_data: Dict) -> BytesIO:
        """
        Generate DOCX from resume data

        Args:
            resume_data: Resume data dictionary

        Returns:
            BytesIO object containing DOCX data
        """
        try:
            document = Document()

            # Set document margins
            sections = document.sections
            for section in sections:
                section.top_margin = Inches(0.5)
                section.bottom_margin = Inches(0.5)
                section.left_margin = Inches(0.75)
                section.right_margin = Inches(0.75)

            # Add personal information
            self._add_personal_info(document, resume_data.get('personal_info', {}))

            # Add professional summary
            if 'professional_summary' in resume_data:
                self._add_section(document, 'Professional Summary', resume_data['professional_summary'])

            # Add work experience
            if 'work_experience' in resume_data:
                self._add_work_experience(document, resume_data['work_experience'])

            # Add education
            if 'education' in resume_data:
                self._add_education(document, resume_data['education'])

            # Add skills
            if 'skills' in resume_data:
                self._add_skills(document, resume_data['skills'])

            # Add projects
            if 'projects' in resume_data:
                self._add_projects(document, resume_data['projects'])

            # Add certifications
            if 'certifications' in resume_data:
                self._add_certifications(document, resume_data['certifications'])

            # Save to BytesIO
            docx_buffer = BytesIO()
            document.save(docx_buffer)
            docx_buffer.seek(0)

            logger.info("DOCX generated successfully")
            return docx_buffer

        except Exception as e:
            logger.error(f"DOCX generation failed: {str(e)}")
            raise

    def _add_personal_info(self, document, personal_info: Dict):
        """Add personal information header"""
        # Name
        name = f"{personal_info.get('first_name', '')} {personal_info.get('last_name', '')}".strip()
        if name:
            heading = document.add_heading(name, level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            heading.runs[0].font.size = Pt(18)
            heading.runs[0].font.color.rgb = RGBColor(0, 0, 0)

        # Contact info
        contact_parts = []
        if personal_info.get('email'):
            contact_parts.append(personal_info['email'])
        if personal_info.get('phone'):
            contact_parts.append(personal_info['phone'])
        if personal_info.get('location'):
            contact_parts.append(personal_info['location'])

        if contact_parts:
            contact = document.add_paragraph(' | '.join(contact_parts))
            contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
            contact.runs[0].font.size = Pt(10)

        # Links
        link_parts = []
        if personal_info.get('linkedin'):
            link_parts.append(personal_info['linkedin'])
        if personal_info.get('github'):
            link_parts.append(personal_info['github'])
        if personal_info.get('portfolio'):
            link_parts.append(personal_info['portfolio'])

        if link_parts:
            links = document.add_paragraph(' | '.join(link_parts))
            links.alignment = WD_ALIGN_PARAGRAPH.CENTER
            links.runs[0].font.size = Pt(9)

        document.add_paragraph()  # Spacing

    def _add_section(self, document, title: str, content: str):
        """Add a section with title and content"""
        heading = document.add_heading(title, level=2)
        heading.runs[0].font.size = Pt(14)
        heading.runs[0].font.color.rgb = RGBColor(0, 0, 0)

        para = document.add_paragraph(content)
        para.runs[0].font.size = Pt(11)

    def _add_work_experience(self, document, experiences: list):
        """Add work experience section"""
        heading = document.add_heading('Work Experience', level=2)
        heading.runs[0].font.size = Pt(14)

        for exp in experiences:
            # Job title and company
            job_heading = document.add_paragraph()
            job_run = job_heading.add_run(f"{exp.get('position', '')} - {exp.get('company', '')}")
            job_run.bold = True
            job_run.font.size = Pt(12)

            # Dates and location
            if exp.get('start_date') or exp.get('end_date'):
                date_para = document.add_paragraph()
                date_text = f"{exp.get('start_date', '')} - {exp.get('end_date', 'Present')}"
                if exp.get('location'):
                    date_text += f" | {exp['location']}"
                date_run = date_para.add_run(date_text)
                date_run.italic = True
                date_run.font.size = Pt(10)

            # Description and bullet points
            if exp.get('description'):
                desc = document.add_paragraph(exp['description'])
                desc.runs[0].font.size = Pt(11)

            if exp.get('bullet_points'):
                for bullet in exp['bullet_points']:
                    bullet_para = document.add_paragraph(bullet, style='List Bullet')
                    bullet_para.runs[0].font.size = Pt(11)

            document.add_paragraph()  # Spacing

    def _add_education(self, document, education_list: list):
        """Add education section"""
        heading = document.add_heading('Education', level=2)
        heading.runs[0].font.size = Pt(14)

        for edu in education_list:
            # Degree and school
            edu_heading = document.add_paragraph()
            edu_run = edu_heading.add_run(f"{edu.get('degree', '')} - {edu.get('school', '')}")
            edu_run.bold = True
            edu_run.font.size = Pt(12)

            # Dates and location
            details = []
            if edu.get('graduation_date'):
                details.append(edu['graduation_date'])
            if edu.get('gpa'):
                details.append(f"GPA: {edu['gpa']}")
            if edu.get('location'):
                details.append(edu['location'])

            if details:
                detail_para = document.add_paragraph(' | '.join(details))
                detail_para.runs[0].italic = True
                detail_para.runs[0].font.size = Pt(10)

            document.add_paragraph()  # Spacing

    def _add_skills(self, document, skills: list):
        """Add skills section"""
        heading = document.add_heading('Skills', level=2)
        heading.runs[0].font.size = Pt(14)

        # Group skills by category if available
        categorized_skills = {}
        for skill in skills:
            category = skill.get('category', 'General')
            if category not in categorized_skills:
                categorized_skills[category] = []
            categorized_skills[category].append(skill.get('name', ''))

        for category, skill_list in categorized_skills.items():
            para = document.add_paragraph()
            if category != 'General':
                cat_run = para.add_run(f"{category}: ")
                cat_run.bold = True
            skills_run = para.add_run(', '.join(skill_list))
            para.runs[0].font.size = Pt(11)

    def _add_projects(self, document, projects: list):
        """Add projects section"""
        heading = document.add_heading('Projects', level=2)
        heading.runs[0].font.size = Pt(14)

        for project in projects:
            # Project name
            proj_heading = document.add_paragraph()
            proj_run = proj_heading.add_run(project.get('name', ''))
            proj_run.bold = True
            proj_run.font.size = Pt(12)

            # Description
            if project.get('description'):
                desc = document.add_paragraph(project['description'])
                desc.runs[0].font.size = Pt(11)

            # Technologies
            if project.get('technologies'):
                tech_para = document.add_paragraph()
                tech_label = tech_para.add_run('Technologies: ')
                tech_label.bold = True
                tech_para.add_run(', '.join(project['technologies']))
                tech_para.runs[0].font.size = Pt(10)

            document.add_paragraph()  # Spacing

    def _add_certifications(self, document, certifications: list):
        """Add certifications section"""
        heading = document.add_heading('Certifications', level=2)
        heading.runs[0].font.size = Pt(14)

        for cert in certifications:
            cert_para = document.add_paragraph()
            cert_run = cert_para.add_run(cert.get('name', ''))
            cert_run.bold = True

            if cert.get('issuer'):
                cert_para.add_run(f" - {cert['issuer']}")

            if cert.get('date'):
                cert_para.add_run(f" ({cert['date']})")

            cert_para.runs[0].font.size = Pt(11)

    def save_docx_file(self, resume, docx_buffer: BytesIO) -> str:
        """
        Save DOCX file to disk/S3 and update resume model

        Args:
            resume: UserResume instance
            docx_buffer: BytesIO containing DOCX data

        Returns:
            File path/URL
        """
        from django.core.files.base import ContentFile

        filename = f"resume_{resume.id}_{resume.user.id}.docx"

        # Save to model's FileField (will handle S3 if configured)
        resume.docx_file.save(
            filename,
            ContentFile(docx_buffer.getvalue()),
            save=True
        )

        logger.info(f"DOCX saved: {resume.docx_file.url}")
        return resume.docx_file.url
